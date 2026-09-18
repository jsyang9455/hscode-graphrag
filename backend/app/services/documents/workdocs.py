"""Extract and analyze broker work documents for GraphRAG learning.

Supports heterogeneous office templates:
- text: .txt .md .csv .json .log
- Word: .docx (OOXML). Legacy .doc is rejected with a clear message.
- PDF: digital text PDFs via pypdf (scanned image PDFs need OCR — not supported yet)
- Excel: .xlsx .xlsm via openpyxl. Legacy .xls is rejected with a clear message.
"""

from __future__ import annotations

import io
import json
import logging
import re
import zipfile
from pathlib import Path
from typing import Any, Optional
from xml.etree import ElementTree as ET

import httpx

from backend.app.core.config import get_settings

logger = logging.getLogger(__name__)

# HSK: 3304 / 3304.99 / 3304.99.10 / 3304.99.1000 / 3304991000
HS_RE = re.compile(r"(?<![\d.])(\d{4}(?:\.\d{2,4}){1,3}|\d{6,10}|\d{4})(?![\d.])")
W_NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".log",
    ".pdf",
    ".docx",
    ".xlsx",
    ".xlsm",
}
UNSUPPORTED_BUT_RELATED = {
    ".doc": "구형 Word(.doc)는 지원하지 않습니다. Word에서 .docx로 저장한 뒤 업로드해 주세요.",
    ".xls": "구형 Excel(.xls)는 지원하지 않습니다. Excel에서 .xlsx로 저장한 뒤 업로드해 주세요.",
    ".ppt": "PowerPoint는 지원하지 않습니다. 텍스트/표가 있는 Word·PDF·Excel로 저장해 주세요.",
    ".pptx": "PowerPoint는 지원하지 않습니다. 텍스트/표가 있는 Word·PDF·Excel로 저장해 주세요.",
}

# Company templates vary widely — match common customs / trade column labels.
FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "product_description": (
        "품명",
        "상품명",
        "품목명",
        "물품명",
        "제품명",
        "상품설명",
        "품목설명",
        "물품개요",
        "description",
        "product",
        "goods",
        "item",
        "article",
    ),
    "material": (
        "재질",
        "소재",
        "성분",
        "구성",
        "원재료",
        "material",
        "composition",
        "fabric",
        "ingredient",
    ),
    "usage": (
        "용도",
        "기능",
        "사용용도",
        "사용처",
        "usage",
        "function",
        "purpose",
        "application",
    ),
    "suggested_hs": (
        "hs",
        "hsk",
        "hs코드",
        "hscode",
        "세번",
        "세번부호",
        "관세율표",
        "분류코드",
        "품목번호",
    ),
    "broker_opinion_excerpt": (
        "의견",
        "사유",
        "근거",
        "판단",
        "분류사유",
        "검토의견",
        "의견서",
        "opinion",
        "reason",
        "rationale",
        "remark",
        "비고",
    ),
}


class ExtractResult:
    def __init__(
        self,
        text: str,
        *,
        format: str,
        ok: bool,
        warning: str | None = None,
        meta: dict[str, Any] | None = None,
    ) -> None:
        self.text = text
        self.format = format
        self.ok = ok
        self.warning = warning
        self.meta = meta or {}

    def as_dict(self) -> dict[str, Any]:
        return {
            "ok": self.ok,
            "format": self.format,
            "warning": self.warning,
            "char_count": len(self.text or ""),
            **self.meta,
        }


def _decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp949", "euc-kr", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _extract_pdf(data: bytes) -> ExtractResult:
    try:
        from pypdf import PdfReader  # type: ignore
    except Exception as exc:  # noqa: BLE001
        return ExtractResult("", format="pdf", ok=False, warning=f"pypdf 미설치: {exc}")

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = reader.pages[:40]
        parts = [(p.extract_text() or "").strip() for p in pages]
        text = "\n\n".join(p for p in parts if p).strip()
        if not text:
            return ExtractResult(
                "",
                format="pdf",
                ok=False,
                warning="PDF에서 텍스트를 찾지 못했습니다. 스캔본(이미지 PDF)은 OCR이 필요합니다. 텍스트 선택 가능한 PDF 또는 Word/Excel로 올려 주세요.",
                meta={"pages": len(reader.pages)},
            )
        return ExtractResult(
            text,
            format="pdf",
            ok=True,
            meta={"pages": len(reader.pages), "pages_read": len(pages)},
        )
    except Exception as exc:  # noqa: BLE001
        logger.warning("pdf extract failed: %s", exc)
        return ExtractResult("", format="pdf", ok=False, warning=f"PDF 추출 실패: {exc}")


def _docx_paragraph_texts(root: ET.Element) -> list[str]:
    texts: list[str] = []
    for p in root.findall(".//w:p", W_NS):
        runs = [n.text for n in p.findall(".//w:t", W_NS) if n.text]
        line = "".join(runs).strip()
        if line:
            texts.append(line)
    return texts


def _docx_table_texts(root: ET.Element) -> list[str]:
    rows_out: list[str] = []
    for tbl in root.findall(".//w:tbl", W_NS):
        for tr in tbl.findall("./w:tr", W_NS):
            cells: list[str] = []
            for tc in tr.findall("./w:tc", W_NS):
                cell_runs = [n.text for n in tc.findall(".//w:t", W_NS) if n.text]
                cell = "".join(cell_runs).strip()
                cells.append(cell)
            if any(cells):
                rows_out.append("\t".join(cells))
    return rows_out


def _extract_docx(data: bytes) -> ExtractResult:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            xml = zf.read("word/document.xml")
            # optional headers/footers — often hold form titles
            extras: list[str] = []
            for name in zf.namelist():
                if re.match(r"word/(header|footer)\d*\.xml$", name):
                    try:
                        hroot = ET.fromstring(zf.read(name))
                        extras.extend(_docx_paragraph_texts(hroot))
                    except Exception:  # noqa: BLE001
                        pass
        root = ET.fromstring(xml)
        paras = _docx_paragraph_texts(root)
        tables = _docx_table_texts(root)
        chunks: list[str] = []
        if extras:
            chunks.append("[머리글/바닥글]\n" + "\n".join(dict.fromkeys(extras)))
        if paras:
            chunks.append("\n".join(paras))
        if tables:
            chunks.append("[표]\n" + "\n".join(tables))
        text = "\n\n".join(chunks).strip()
        if not text:
            return ExtractResult(
                "",
                format="docx",
                ok=False,
                warning="Word 문서에서 텍스트를 찾지 못했습니다.",
            )
        return ExtractResult(
            text,
            format="docx",
            ok=True,
            meta={"paragraphs": len(paras), "table_rows": len(tables)},
        )
    except Exception as exc:  # noqa: BLE001
        logger.warning("docx extract failed: %s", exc)
        return ExtractResult("", format="docx", ok=False, warning=f"Word 추출 실패: {exc}")


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _extract_xlsx(data: bytes, filename: str) -> ExtractResult:
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception as exc:  # noqa: BLE001
        return ExtractResult("", format="xlsx", ok=False, warning=f"openpyxl 미설치: {exc}")

    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        sheet_names = list(wb.sheetnames)
        sheet_chunks: list[str] = []
        rows_total = 0
        for sheet in wb.worksheets:
            rows: list[str] = []
            for i, row in enumerate(sheet.iter_rows(values_only=True)):
                if i >= 500:  # guard huge sheets
                    rows.append("… (행 한도 초과, 일부만 추출)")
                    break
                vals = [_cell_str(c) for c in row]
                if not any(vals):
                    continue
                # trim trailing empties
                while vals and not vals[-1]:
                    vals.pop()
                rows.append("\t".join(vals))
                rows_total += 1
            if rows:
                sheet_chunks.append(f"[시트: {sheet.title}]\n" + "\n".join(rows))
        wb.close()
        text = "\n\n".join(sheet_chunks).strip()
        if not text:
            return ExtractResult(
                "",
                format="xlsx",
                ok=False,
                warning="Excel 시트에서 값을 찾지 못했습니다.",
                meta={"sheets": len(sheet_names)},
            )
        return ExtractResult(
            text,
            format="xlsx",
            ok=True,
            meta={"sheets": len(sheet_chunks), "rows": rows_total, "filename": filename},
        )
    except Exception as exc:  # noqa: BLE001
        logger.warning("xlsx extract failed: %s", exc)
        return ExtractResult("", format="xlsx", ok=False, warning=f"Excel 추출 실패: {exc}")


def extract_document(filename: str, data: bytes, mime: str = "") -> ExtractResult:
    name = (filename or "upload.bin").lower()
    suffix = Path(name).suffix
    mime = (mime or "").lower()

    if suffix in UNSUPPORTED_BUT_RELATED:
        return ExtractResult("", format=suffix.lstrip("."), ok=False, warning=UNSUPPORTED_BUT_RELATED[suffix])

    if suffix in {".txt", ".md", ".csv", ".json", ".log"} or mime.startswith("text/"):
        return ExtractResult(_decode_text(data), format=suffix.lstrip(".") or "text", ok=True)

    if suffix == ".pdf" or mime == "application/pdf":
        return _extract_pdf(data)

    if suffix == ".docx" or "wordprocessingml" in mime:
        return _extract_docx(data)

    if suffix in {".xlsx", ".xlsm"} or "spreadsheetml" in mime:
        return _extract_xlsx(data, name)

    if suffix == ".doc" or mime == "application/msword":
        return ExtractResult("", format="doc", ok=False, warning=UNSUPPORTED_BUT_RELATED[".doc"])

    if suffix == ".xls" or mime == "application/vnd.ms-excel":
        return ExtractResult("", format="xls", ok=False, warning=UNSUPPORTED_BUT_RELATED[".xls"])

    # unknown binary — do not dump garbage as "text"
    return ExtractResult(
        "",
        format=suffix.lstrip(".") or "unknown",
        ok=False,
        warning=(
            f"지원하지 않는 형식({suffix or mime or 'unknown'})입니다. "
            f"지원: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        ),
    )


def extract_text_from_bytes(filename: str, data: bytes, mime: str = "") -> str:
    """Backward-compatible text-only API used by routes."""
    result = extract_document(filename, data, mime)
    if result.ok and result.text.strip():
        return result.text
    if result.warning:
        return f"(텍스트 추출 실패) {result.warning}"
    return f"(텍스트 추출 실패) 파일명: {filename}"


def _normalize_header(cell: str) -> str:
    return re.sub(r"[\s_\-:/：]+", "", (cell or "").lower())


def _match_field(header: str) -> Optional[str]:
    h = _normalize_header(header)
    if not h:
        return None
    for field, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            a = _normalize_header(alias)
            if a and (a in h or h in a):
                return field
    return None


def _parse_tabular_kv(text: str) -> dict[str, str]:
    """Pull field values from TSV/CSV-like rows and '라벨: 값' lines across templates."""
    found: dict[str, str] = {}
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]

    # header + data row pattern (common in Excel exports)
    for i, ln in enumerate(lines[:80]):
        if "\t" in ln or ("," in ln and ln.count(",") >= 2):
            sep = "\t" if "\t" in ln else ","
            headers = [c.strip() for c in ln.split(sep)]
            mapped = [(idx, _match_field(h)) for idx, h in enumerate(headers)]
            if sum(1 for _, f in mapped if f) >= 2 and i + 1 < len(lines):
                data_cells = [c.strip() for c in lines[i + 1].split(sep)]
                for idx, field in mapped:
                    if not field or field in found:
                        continue
                    if idx < len(data_cells) and data_cells[idx]:
                        found[field] = data_cells[idx][:500]
                if found:
                    break

    # label: value / label\tvalue rows
    for ln in lines[:120]:
        if ":" in ln or "：" in ln or "\t" in ln:
            if "\t" in ln:
                parts = ln.split("\t", 1)
            else:
                parts = re.split(r"[:：]", ln, maxsplit=1)
            if len(parts) != 2:
                continue
            field = _match_field(parts[0])
            val = parts[1].strip()
            if field and val and field not in found:
                found[field] = val[:500]
    return found


def _heuristic_analysis(text: str, doc_type: str) -> dict[str, Any]:
    extract_failed = text.startswith("(텍스트 추출 실패)")
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    hs_codes: list[str] = []
    for m in HS_RE.finditer(text):
        code = m.group(1)
        if code not in hs_codes:
            hs_codes.append(code)

    kv = _parse_tabular_kv(text)

    description = kv.get("product_description") or ""
    if not description:
        for ln in lines[:60]:
            if ln.startswith("[시트:") or ln.startswith("[표]") or ln.startswith("[머리글"):
                continue
            if ln.startswith("(텍스트 추출 실패)"):
                continue
            if HS_RE.fullmatch(ln.replace(" ", "")):
                continue
            if _match_field(ln.split("\t")[0] if "\t" in ln else ln.split(":")[0]):
                continue
            if len(ln) >= 6 and not ln.lower().startswith(("http", "www")):
                description = ln[:300]
                break
    if not description and lines:
        description = lines[0][:300]

    material = kv.get("material")
    usage = kv.get("usage")
    suggested = kv.get("suggested_hs")
    if suggested:
        m = HS_RE.search(suggested)
        if m and m.group(1) not in hs_codes:
            hs_codes.insert(0, m.group(1))
        elif m:
            hs_codes = [m.group(1)] + [c for c in hs_codes if c != m.group(1)]

    opinion = kv.get("broker_opinion_excerpt") or ""
    if not opinion:
        for key in ("의견", "사유", "근거", "판단", "opinion", "reason"):
            for ln in lines:
                if key in ln.lower() or key in ln:
                    opinion = ln[:500]
                    break
            if opinion:
                break
    if not opinion:
        opinion = "\n".join(lines[:8])[:800]

    summary = (
        f"업무자료에서 상품 '{description or '(미추출)'}' "
        f"및 HS 후보 {', '.join(hs_codes[:3]) or '없음'}을(를) 추출했습니다."
    )
    if extract_failed:
        summary = text[:240]

    return {
        "source": "heuristic",
        "doc_type": doc_type,
        "product_description": description,
        "material": material,
        "usage": usage,
        "suggested_hs": hs_codes[0] if hs_codes else None,
        "candidate_hs": hs_codes[:5],
        "broker_opinion_excerpt": opinion,
        "keywords": list(
            dict.fromkeys(re.findall(r"[A-Za-z가-힣]{2,}", (description or "").lower()))
        )[:12],
        "summary_ko": summary,
        "extract_ok": not extract_failed,
        "matched_fields": sorted(kv.keys()),
    }


def _llm_analysis(text: str, doc_type: str, draft: dict[str, Any]) -> dict[str, Any]:
    settings = get_settings()
    if not settings.use_llm or not settings.openai_api_key:
        return draft
    if text.startswith("(텍스트 추출 실패)"):
        return draft
    payload = {
        "doc_type": doc_type,
        "text_excerpt": text[:8000],
        "draft": draft,
        "instruction": (
            "회사마다 양식이 다른 관세사 업무자료(의견서/인보이스/엑셀표 등)입니다. "
            "JSON만 반환. 키: product_description, material, usage, suggested_hs, "
            "candidate_hs(list), broker_opinion_excerpt, keywords(list), summary_ko. "
            "없는 값은 null. HS는 문서에 명시된 코드만. 표 헤더가 달라도 의미로 매핑하세요."
        ),
    }
    try:
        with httpx.Client(timeout=35.0) as client:
            resp = client.post(
                "https://api.openai.com/v1/chat/completions",
                headers={
                    "Authorization": f"Bearer {settings.openai_api_key}",
                    "Content-Type": "application/json",
                },
                json={
                    "model": settings.openai_model,
                    "temperature": 0.1,
                    "response_format": {"type": "json_object"},
                    "messages": [
                        {
                            "role": "system",
                            "content": (
                                "당신은 한국 관세사 업무자료 분석 보조원입니다. "
                                "근거 없는 HS를 만들지 마세요. 양식이 달라도 핵심 필드만 추출합니다."
                            ),
                        },
                        {"role": "user", "content": json.dumps(payload, ensure_ascii=False)},
                    ],
                },
            )
            resp.raise_for_status()
            data = json.loads(resp.json()["choices"][0]["message"]["content"])
    except Exception as exc:  # noqa: BLE001
        logger.warning("document LLM analysis failed: %s", exc)
        draft["llm_error"] = str(exc)[:200]
        return draft

    out = dict(draft)
    for k in (
        "product_description",
        "material",
        "usage",
        "suggested_hs",
        "broker_opinion_excerpt",
        "summary_ko",
    ):
        if data.get(k) is not None:
            out[k] = data[k]
    if isinstance(data.get("candidate_hs"), list):
        out["candidate_hs"] = data["candidate_hs"][:8]
    if isinstance(data.get("keywords"), list):
        out["keywords"] = [str(x) for x in data["keywords"][:12]]
    out["source"] = "llm"
    return out


def analyze_document_text(text: str, doc_type: str = "opinion") -> dict[str, Any]:
    draft = _heuristic_analysis(text or "", doc_type)
    return _llm_analysis(text or "", doc_type, draft)


def ensure_upload_dir() -> Path:
    path = Path("data/runtime/uploads")
    path.mkdir(parents=True, exist_ok=True)
    return path
